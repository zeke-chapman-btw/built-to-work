import csv
import io
import tempfile
import uuid
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.events.models import Event, EventRegistration
from apps.participants.historical import (create_batch, export_participants, import_ready,
                                          preview_batch, resolve_review)
from apps.participants.models import (ConsentAcceptance, HistoricalEventRecord,
                                      HistoricalExportRecord, HistoricalImportBatch,
                                      HistoricalImportRow, HistoricalMappingTemplate,
                                      HistoricalSimulatorScore, Participant)


class HistoricalDataTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        setting = override_settings(PRIVATE_HISTORICAL_ROOT=self.temp.name)
        setting.enable()
        self.addCleanup(setting.disable)
        self.actor = get_user_model().objects.create_superuser(
            username="history-admin", email="history@example.test", password="test-password")
        self.other = get_user_model().objects.create_user(
            username="history-visitor", password="test-password")
        self.mapping = {"first_name": "First", "last_name": "Last", "email": "Email"}

    def batch(self, text, name="history.csv"):
        file = SimpleUploadedFile(name, text.encode("utf-8"), content_type="text/csv")
        return create_batch(uploaded_file=file, actor=self.actor)

    def test_csv_preview_import_without_event_or_ticket(self):
        batch = self.batch("First,Last,Email\nAva,Stone,AVA@example.test\n")
        rows = list(preview_batch(batch, self.mapping, actor=self.actor))
        self.assertEqual(rows[0].status, HistoricalImportRow.Status.READY)
        self.assertEqual(import_ready(batch, actor=self.actor), 1)
        person = Participant.objects.get(contact_email="ava@example.test")
        self.assertEqual(str(rows[0].normalized_data["first_name"]), "Ava")
        self.assertFalse(EventRegistration.objects.filter(participant=person).exists())
        self.assertFalse(ConsentAcceptance.objects.filter(participant=person).exists())
        self.assertEqual(HistoricalEventRecord.objects.count(), 0)
        self.assertEqual(HistoricalImportBatch.objects.count(), 1)
        self.assertEqual(import_ready.__name__, "import_ready")
        with self.assertRaises(ValidationError):
            import_ready(batch, actor=self.actor)

    def test_xlsx_historical_event_score_and_original_uuid(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["First", "Last", "UUID", "Event", "Date", "Group", "Score", "Method", "Source", "Scored"])
        identifier = uuid.uuid4()
        sheet.append(["Maya", "Field", str(identifier), "Legacy Expo", "2023-05-01", "Class A",
                      "87.5", "original system", "legacy export", "2023-05-01T14:30:00+00:00"])
        stream = io.BytesIO()
        book.save(stream)
        batch = create_batch(uploaded_file=SimpleUploadedFile("legacy.xlsx", stream.getvalue()), actor=self.actor)
        mapping = {"first_name": "First", "last_name": "Last", "identity_uuid": "UUID",
                   "event_name": "Event", "event_date": "Date", "group_name": "Group",
                   "simulator_score": "Score", "score_method": "Method", "score_source": "Source",
                   "scored_at": "Scored"}
        self.assertEqual(list(preview_batch(batch, mapping, actor=self.actor))[0].status, "ready")
        self.assertEqual(import_ready(batch, actor=self.actor), 1)
        person = Participant.objects.get(identity_uuid=identifier)
        event = HistoricalEventRecord.objects.get(participant=person)
        score = HistoricalSimulatorScore.objects.get(participant=person)
        self.assertEqual(event.group_name, "Class A")
        self.assertEqual(score.score_method, "original system")
        self.assertEqual(score.scored_at.hour, 14)
        self.assertEqual(EventRegistration.objects.count(), 0)

    def test_row_errors_and_duplicate_review_are_visible(self):
        batch = self.batch("First,Last,Email\nNo,Mail,\nNo,Mail,\nToo,Many,x@y.test,extra\n,Missing,x@y.test\n")
        rows = list(preview_batch(batch, self.mapping, actor=self.actor))
        self.assertEqual([r.status for r in rows], ["ready", "ready", "error", "error"])
        self.assertIn("Column count", rows[2].error)
        self.assertEqual(import_ready(batch, actor=self.actor), 2)
        self.assertEqual(Participant.objects.count(), 2)
        batch2 = self.batch("First,Last,Email\nA,One,same@example.test\nB,Two,same@example.test\n")
        rows = list(preview_batch(batch2, self.mapping, actor=self.actor))
        self.assertEqual([r.status for r in rows], ["ready", "review"])
        self.assertEqual(import_ready(batch2, actor=self.actor), 1)
        with self.assertRaises(ValidationError):
            resolve_review(rows[1], actor=self.actor, decision="separate", reason="")
        resolved = resolve_review(rows[1], actor=self.actor, decision="separate", reason="Verified separate person")
        self.assertEqual(resolved.status, "imported")
        with self.assertRaises(ValidationError):
            resolve_review(rows[1], actor=self.actor, decision="separate", reason="repeat")

    def test_existing_participant_requires_review_and_is_not_overwritten(self):
        person = Participant.objects.create(first_name="Existing", last_name="Person",
                                            contact_email="same@example.test")
        batch = self.batch("First,Last,Email\nDifferent,Name,same@example.test\n")
        row = list(preview_batch(batch, self.mapping, actor=self.actor))[0]
        self.assertEqual(row.status, "review")
        self.assertEqual(import_ready(batch, actor=self.actor), 0)
        self.assertEqual(Participant.objects.count(), 1)
        resolve_review(row, actor=self.actor, decision="link", reason="Verified identity", participant=person)
        person.refresh_from_db()
        self.assertEqual(person.first_name, "Existing")
        self.assertEqual(Participant.objects.count(), 1)

    def test_export_filters_and_escapes_formulas_in_csv_and_excel(self):
        person = Participant.objects.create(first_name="=danger", last_name="Example")
        other = Participant.objects.create(first_name="Other", last_name="No event")
        batch = self.batch("First,Last,Email\nA,B,a@example.test\n")
        HistoricalEventRecord.objects.create(participant=person, event_name="Legacy", group_name="G1",
                                             source_batch=batch, source_row_number=2)
        csv_payload, count = export_participants(fields=["first_name", "event_name", "group_name"],
                            filters={"group": "G1"}, file_format="csv", actor=self.actor)
        self.assertEqual(count, 1)
        rows = list(csv.reader(io.StringIO(csv_payload.decode("utf-8-sig"))))
        self.assertEqual(rows[1], ["'=danger", "Legacy", "G1"])
        xlsx_payload, count = export_participants(fields=["first_name", "event_name"],
                            filters={"without_events": True}, file_format="xlsx", actor=self.actor)
        self.assertEqual(count, 1)
        sheet = load_workbook(io.BytesIO(xlsx_payload)).active
        self.assertEqual(sheet.cell(2, 1).value, other.first_name)
        self.assertEqual(HistoricalExportRecord.objects.count(), 2)

    def test_staff_routes_reject_unauthorized_post_and_record_audit(self):
        batch = self.batch("First,Last,Email\nA,B,a@example.test\n")
        map_url = reverse("participants:historical-map", kwargs={"batch_id": batch.pk})
        self.client.force_login(self.other, backend="django.contrib.auth.backends.ModelBackend")
        r = self.client.post(map_url, {"map_first_name": "First"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.client.post(reverse("participants:historical-export"),
                                          {"fields": ["first_name"]}).status_code, 403)
        self.client.force_login(self.actor, backend="django.contrib.auth.backends.ModelBackend")
        response = self.client.get(reverse("participants:historical-index"))
        self.assertEqual(response.status_code, 200)
        response = self.client.post(map_url, {"map_first_name": "First", "map_last_name": "Last",
                                              "map_email": "Email", "template_name": "Legacy mapping"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(HistoricalMappingTemplate.objects.count(), 1)
        response = self.client.post(reverse("participants:historical-batch", kwargs={"batch_id": batch.pk}))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Participant.objects.count(), 1)
        response = self.client.post(reverse("participants:historical-export"),
                                    {"fields": ["first_name"], "format": "csv", "template_name": "Names"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(HistoricalExportRecord.objects.count(), 1)

    def test_import_does_not_change_existing_registration(self):
        now = timezone.now()
        event = Event.objects.create(name="Existing Event", start_at=now, end_at=now + timedelta(hours=2))
        person = Participant.objects.create(first_name="Registered", last_name="Person")
        registration = EventRegistration.objects.create(event=event, participant=person)
        batch = self.batch("First,Last,Email\nHistory,Only,new@example.test\n")
        preview_batch(batch, self.mapping, actor=self.actor)
        import_ready(batch, actor=self.actor)
        registration.refresh_from_db()
        self.assertEqual(registration.participant_id, person.pk)
        self.assertEqual(EventRegistration.objects.count(), 1)

    def test_name_and_birth_date_is_flagged_for_review(self):
        Participant.objects.create(first_name="Lee", last_name="Morgan", date_of_birth="1991-06-03")
        batch = self.batch("First,Last,Birth\nLee,Morgan,1991-06-03\n")
        rows = list(preview_batch(batch, {"first_name": "First", "last_name": "Last",
                                          "date_of_birth": "Birth"}, actor=self.actor))
        self.assertEqual(rows[0].status, "review")
        self.assertEqual(import_ready(batch, actor=self.actor), 0)
        self.assertEqual(Participant.objects.count(), 1)

    def test_score_export_stays_with_its_original_event(self):
        person = Participant.objects.create(first_name="History", last_name="Scores")
        batch = self.batch("First,Last\nHistory,Scores\n")
        first = HistoricalEventRecord.objects.create(participant=person, event_name="Event A",
                    source_batch=batch, source_row_number=2)
        HistoricalEventRecord.objects.create(participant=person, event_name="Event B",
                    source_batch=batch, source_row_number=3)
        HistoricalSimulatorScore.objects.create(participant=person, event_record=first, raw_score="77",
                    score_method="legacy", source="export 2022", source_batch=batch, source_row_number=2)
        payload, count = export_participants(fields=["event_name", "simulator_score", "score_method"],
                    filters={}, file_format="csv", actor=self.actor)
        self.assertEqual(count, 2)
        rows = list(csv.reader(io.StringIO(payload.decode("utf-8-sig"))))
        self.assertEqual(rows[1][0], "Event A")
        self.assertEqual(rows[1][1], "77.0000")
        self.assertEqual(rows[2], ["Event B", "", ""])

    def test_staff_pages_render_and_uploaded_sources_are_private(self):
        batch = self.batch("First,Last\nA,B\n")
        self.assertFalse(batch.source_file.storage.base_url)
        self.client.force_login(self.actor, backend="django.contrib.auth.backends.ModelBackend")
        urls = [reverse("participants:historical-index"),
                reverse("participants:historical-map", kwargs={"batch_id": batch.pk}),
                reverse("participants:historical-batch", kwargs={"batch_id": batch.pk}),
                reverse("participants:historical-export")]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_duplicate_review_is_permissioned_and_audited(self):
        from apps.core.models import AuditLog
        person = Participant.objects.create(first_name="Same", last_name="Person",
                                            contact_email="same@example.test")
        batch = self.batch("First,Last,Email\nSame,Person,same@example.test\n")
        row = list(preview_batch(batch, self.mapping, actor=self.actor))[0]
        url = reverse("participants:historical-review", kwargs={"row_id": row.pk})
        self.client.force_login(self.other, backend="django.contrib.auth.backends.ModelBackend")
        self.assertEqual(self.client.post(url, {"decision": "link", "reason": "unsafe",
                                                "participant_uuid": str(person.identity_uuid)}).status_code, 403)
        self.assertEqual(row.status, "review")
        self.client.force_login(self.actor, backend="django.contrib.auth.backends.ModelBackend")
        response = self.client.post(url, {"decision": "link", "reason": "Checked source records",
                                          "participant_uuid": str(person.identity_uuid)})
        self.assertEqual(response.status_code, 302)
        row.refresh_from_db()
        self.assertEqual(row.participant_id, person.pk)
        self.assertTrue(AuditLog.objects.filter(action="historical.import.resolve").exists())
        self.assertEqual(Participant.objects.count(), 1)
