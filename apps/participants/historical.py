"""Staff-controlled historical spreadsheet import and export. No tickets or consent are created."""
import csv
import hashlib
import io
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from apps.events.models import Event
from apps.events.services import audit_action
from .models import (HistoricalEventRecord, HistoricalImportBatch, HistoricalImportRow,
                     HistoricalSimulatorScore, Participant)
from .normalization import normalize_email, normalize_phone

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_XLSX_EXPANDED = 30 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 100
IMPORT_FIELDS = ("identity_uuid", "first_name", "last_name", "preferred_name", "email", "phone",
                 "date_of_birth", "city", "state", "postal_code", "event_uuid", "event_name",
                 "event_date", "group_name", "simulator_score", "score_method", "score_source", "scored_at")
EXPORT_FIELDS = ("identity_uuid", "first_name", "last_name", "preferred_name", "email", "phone",
                 "date_of_birth", "city", "state", "postal_code", "event_name", "event_date",
                 "group_name", "registration_source", "simulator_score", "score_method", "score_source")


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _read_rows(batch):
    with batch.source_file.open("rb") as stream:
        data = stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValidationError("Spreadsheet exceeds the 5 MB limit.")
    suffix = Path(batch.original_filename).suffix.lower()
    if suffix == ".csv":
        try:
            reader = csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""))
            table = list(reader)
        except (UnicodeError, csv.Error) as exc:
            raise ValidationError("Invalid UTF-8 CSV file.") from exc
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 250 or sum(item.file_size for item in archive.infolist()) > MAX_XLSX_EXPANDED:
                    raise ValidationError("Excel workbook exceeds the safe expanded-size limit.")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                table = [[_cell(value) for value in row] for row in workbook.active.iter_rows(values_only=True)]
            finally:
                workbook.close()
        except (zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
            raise ValidationError("Invalid Excel workbook.") from exc
    else:
        raise ValidationError("Upload a CSV or .xlsx file.")
    if not table:
        raise ValidationError("Spreadsheet is empty.")
    if len(table) - 1 > MAX_ROWS or max(map(len, table)) > MAX_COLUMNS:
        raise ValidationError("Spreadsheet exceeds the row or column limit.")
    headers = [_cell(value) for value in table[0]]
    if not all(headers) or len(headers) != len(set(headers)):
        raise ValidationError("Headers must be present and unique.")
    rows = []
    for row_number, values in enumerate(table[1:], 2):
        if any(_cell(value) for value in values):
            raw = dict(zip(headers, (_cell(value) for value in values)))
            if len(values) != len(headers):
                raw["__row_error__"] = "Column count does not match the header."
            rows.append((row_number, raw))
    return headers, rows


def create_batch(*, uploaded_file, actor):
    name = Path(uploaded_file.name).name[:255]
    if Path(name).suffix.lower() not in (".csv", ".xlsx") or uploaded_file.size > MAX_FILE_BYTES:
        raise ValidationError("Upload a CSV or .xlsx file up to 5 MB.")
    digest = hashlib.sha256()
    for chunk in uploaded_file.chunks():
        digest.update(chunk)
    uploaded_file.seek(0)
    batch = HistoricalImportBatch.objects.create(source_file=uploaded_file, original_filename=name,
                                                  source_sha256=digest.hexdigest(), created_by=actor)
    audit_action(actor=actor, action="historical.import.upload", instance=batch,
                 new_data={"sha256": batch.source_sha256})
    return batch


def _date(value, field):
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except (ValueError, TypeError) as exc:
        raise ValidationError(f"{field} must be an ISO date (YYYY-MM-DD).") from exc


def _score_timestamp(value):
    try:
        parsed = datetime.fromisoformat(value)
        return timezone.make_aware(parsed) if timezone.is_naive(parsed) else parsed
    except (ValueError, TypeError) as exc:
        raise ValidationError("scored_at must be an ISO date and time.") from exc


def _normalize(raw, mapping):
    if raw.get("__row_error__"):
        raise ValidationError(raw["__row_error__"])
    data = {field: _cell(raw.get(column, "")) for field, column in mapping.items() if column}
    if not data.get("first_name") or not data.get("last_name"):
        raise ValidationError("First and last names are required.")
    for field, value in data.items():
        if value.startswith("="):
            raise ValidationError("Spreadsheet formulas are not accepted as participant data.")
    data["email"] = normalize_email(data.get("email", ""))
    data["phone"] = normalize_phone(data.get("phone", ""))
    if data.get("identity_uuid"):
        try:
            data["identity_uuid"] = str(uuid.UUID(data["identity_uuid"]))
        except ValueError as exc:
            raise ValidationError("Invalid participant UUID.") from exc
    if data.get("event_uuid"):
        try:
            data["event_uuid"] = str(uuid.UUID(data["event_uuid"]))
        except ValueError as exc:
            raise ValidationError("Invalid Event UUID.") from exc
        if not Event.objects.filter(pk=data["event_uuid"]).exists():
            raise ValidationError("Event UUID does not exist.")
    for field in ("date_of_birth", "event_date"):
        if data.get(field):
            data[field] = _date(data[field], field).isoformat()
    if data.get("simulator_score"):
        try:
            score = Decimal(data["simulator_score"])
            if not score.is_finite() or score.as_tuple().exponent < -4 or score.adjusted() >= 10:
                raise InvalidOperation
            data["simulator_score"] = str(score)
        except InvalidOperation as exc:
            raise ValidationError("Invalid simulator score.") from exc
        if not data.get("score_method") or not data.get("score_source"):
            raise ValidationError("Simulator score requires its original method and source.")
    if data.get("scored_at"):
        data["scored_at"] = _score_timestamp(data["scored_at"]).isoformat()
    return data


def _matches(data):
    query = Q(pk__in=[])
    if data.get("identity_uuid"):
        query |= Q(identity_uuid=data["identity_uuid"])
    if data.get("email"):
        query |= Q(contact_email__iexact=data["email"])
    if data.get("phone"):
        query |= Q(contact_phone=data["phone"])
    if data.get("first_name") and data.get("last_name") and data.get("date_of_birth"):
        query |= Q(first_name__iexact=data["first_name"], last_name__iexact=data["last_name"],
                   date_of_birth=data["date_of_birth"])
    if not any(data.get(field) for field in ("identity_uuid", "email", "phone", "date_of_birth")):
        return []
    return [str(value) for value in Participant.objects.filter(query, kind=Participant.Kind.PERSON)
            .values_list("identity_uuid", flat=True)[:10]]


def preview_batch(batch, mapping, *, actor):
    if batch.status == HistoricalImportBatch.Status.IMPORTED:
        raise ValidationError("An imported batch cannot be remapped.")
    if set(mapping) - set(IMPORT_FIELDS) or not mapping.get("first_name") or not mapping.get("last_name"):
        raise ValidationError("Map first name and last name to spreadsheet columns.")
    headers, source_rows = _read_rows(batch)
    if any(column not in headers for column in mapping.values() if column):
        raise ValidationError("A mapped column is missing from this spreadsheet.")
    seen = set()
    with transaction.atomic():
        batch.rows.all().delete()
        for number, raw in source_rows:
            try:
                data = _normalize(raw, mapping)
                matches = _matches(data)
                key = tuple(data.get(field) or "" for field in ("identity_uuid", "email", "phone"))
                secondary = (data.get("first_name", "").casefold(), data.get("last_name", "").casefold(),
                             data.get("date_of_birth", "")) if data.get("date_of_birth") else None
                repeated = (any(key) and key in seen) or (secondary is not None and secondary in seen)
                if any(key):
                    seen.add(key)
                if secondary is not None:
                    seen.add(secondary)
                status = HistoricalImportRow.Status.REVIEW if matches or repeated else HistoricalImportRow.Status.READY
                error = "Possible duplicate in this file." if repeated else ""
            except ValidationError as exc:
                data, matches, status, error = {}, [], HistoricalImportRow.Status.ERROR, "; ".join(exc.messages)
            HistoricalImportRow.objects.create(batch=batch, row_number=number, raw_data=raw,
                normalized_data=data, status=status, error=error, possible_matches=matches)
        batch.mapping = mapping
        batch.status = HistoricalImportBatch.Status.PREVIEWED
        batch.save(update_fields=("mapping", "status"))
        audit_action(actor=actor, action="historical.import.preview", instance=batch,
                     new_data={"rows": len(source_rows)})
    return batch.rows.all()


def _import_row(row, *, actor, participant=None):
    data = row.normalized_data
    if participant is None:
        fields = {"first_name": data["first_name"], "last_name": data["last_name"],
                  "preferred_name": data.get("preferred_name", ""), "contact_email": data.get("email", ""),
                  "contact_phone": data.get("phone", ""), "city": data.get("city", ""),
                  "state": data.get("state", ""), "postal_code": data.get("postal_code", "")}
        if data.get("date_of_birth"):
            fields["date_of_birth"] = _date(data["date_of_birth"], "date_of_birth")
        if data.get("identity_uuid"):
            fields["identity_uuid"] = uuid.UUID(data["identity_uuid"])
        participant = Participant.objects.create(**fields)
    event_record = None
    if any(data.get(field) for field in ("event_name", "event_date", "event_uuid", "group_name")):
        event = Event.objects.filter(pk=data["event_uuid"]).first() if data.get("event_uuid") else None
        event_record, _ = HistoricalEventRecord.objects.get_or_create(source_batch=row.batch,
            source_row_number=row.row_number, defaults={"participant": participant, "event": event,
            "event_name": data.get("event_name") or (event.name if event else "Historical Event"),
            "event_date": _date(data.get("event_date"), "event_date"), "group_name": data.get("group_name", "")})
    if data.get("simulator_score"):
        HistoricalSimulatorScore.objects.get_or_create(source_batch=row.batch, source_row_number=row.row_number,
            defaults={"participant": participant, "event_record": event_record,
                      "raw_score": Decimal(data["simulator_score"]), "score_method": data["score_method"],
                      "source": data["score_source"], "scored_at": _score_timestamp(data["scored_at"]) if data.get("scored_at") else None})
    row.participant = participant
    row.status = HistoricalImportRow.Status.IMPORTED
    row.save(update_fields=("participant", "status"))
    audit_action(actor=actor, action="historical.import.row", instance=row,
                 new_data={"participant_uuid": str(participant.identity_uuid), "row": row.row_number})
    return participant


def import_ready(batch, *, actor):
    if batch.status != HistoricalImportBatch.Status.PREVIEWED:
        raise ValidationError("Preview and map the batch before import.")
    imported = 0
    for row in batch.rows.filter(status=HistoricalImportRow.Status.READY):
        try:
            with transaction.atomic():
                if _matches(row.normalized_data):
                    row.status = HistoricalImportRow.Status.REVIEW
                    row.error = "Possible participant match appeared after preview."
                    row.save(update_fields=("status", "error"))
                else:
                    _import_row(row, actor=actor)
                    imported += 1
        except (ValidationError, IntegrityError, ValueError) as exc:
            row.status = HistoricalImportRow.Status.ERROR
            row.error = str(exc)[:1000]
            row.save(update_fields=("status", "error"))
    batch.status = HistoricalImportBatch.Status.IMPORTED
    batch.completed_at = timezone.now()
    batch.save(update_fields=("status", "completed_at"))
    audit_action(actor=actor, action="historical.import.complete", instance=batch,
                 new_data={"imported": imported, "review": batch.rows.filter(status="review").count(),
                           "errors": batch.rows.filter(status="error").count()})
    return imported


def resolve_review(row, *, actor, decision, reason, participant=None):
    if not reason.strip() or decision not in ("link", "separate"):
        raise ValidationError("A decision and audit reason are required.")
    with transaction.atomic():
        row = HistoricalImportRow.objects.select_for_update().get(pk=row.pk)
        if row.status != HistoricalImportRow.Status.REVIEW:
            raise ValidationError("This row is no longer awaiting review.")
        if decision == "link":
            if participant is None or participant.kind != Participant.Kind.PERSON:
                raise ValidationError("Choose a person to link.")
            source_uuid = row.normalized_data.get("identity_uuid")
            if source_uuid and str(participant.identity_uuid) != source_uuid:
                raise ValidationError("A different source UUID cannot be silently replaced.")
        else:
            participant = None
            if row.normalized_data.get("identity_uuid") and Participant.objects.filter(
                    identity_uuid=row.normalized_data["identity_uuid"]).exists():
                raise ValidationError("Source UUID already belongs to an existing participant.")
        _import_row(row, actor=actor, participant=participant)
        row.review_reason = reason.strip()
        row.reviewed_by = actor
        row.reviewed_at = timezone.now()
        row.save(update_fields=("review_reason", "reviewed_by", "reviewed_at"))
        audit_action(actor=actor, action="historical.import.resolve", instance=row,
                     new_data={"decision": decision, "participant_uuid": str(row.participant.identity_uuid)},
                     reason=reason.strip())
        return row


def safe_spreadsheet_cell(value):
    """Prevent CSV/Excel clients from evaluating imported text as a formula."""
    text = "" if value is None else str(value)
    if text and text.lstrip(" ").startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + text
    return text


def export_participants(*, fields, filters, file_format, actor):
    if file_format not in ("csv", "xlsx") or not fields or set(fields) - set(EXPORT_FIELDS):
        raise ValidationError("Choose supported fields and CSV or Excel format.")
    people = Participant.objects.filter(kind=Participant.Kind.PERSON)
    if filters.get("without_events"):
        people = people.filter(event_registrations__isnull=True, historical_events__isnull=True)
    if filters.get("event_id"):
        try:
            event_id = uuid.UUID(filters["event_id"])
        except ValueError as exc:
            raise ValidationError("Invalid Event filter.") from exc
        people = people.filter(Q(event_registrations__event_id=event_id) | Q(historical_events__event_id=event_id))
    if filters.get("group"):
        people = people.filter(Q(event_registrations__group__name__iexact=filters["group"]) |
                               Q(historical_events__group_name__iexact=filters["group"]))
    if filters.get("registration_source"):
        source = filters["registration_source"]
        if source == "historical":
            people = people.filter(historical_events__isnull=False)
        else:
            people = people.filter(event_registrations__source=source)
    start_date = _date(filters.get("date_from"), "date_from")
    end_date = _date(filters.get("date_to"), "date_to")
    if start_date and end_date and start_date > end_date:
        raise ValidationError("Start date must not follow end date.")
    if start_date:
        people = people.filter(Q(event_registrations__registered_at__date__gte=start_date) |
                               Q(historical_events__event_date__gte=start_date))
    if end_date:
        people = people.filter(Q(event_registrations__registered_at__date__lte=end_date) |
                               Q(historical_events__event_date__lte=end_date))
    people = people.distinct().prefetch_related("event_registrations__event", "event_registrations__group",
                                                "historical_events", "historical_simulator_scores")
    rows = []
    for person in people.iterator(chunk_size=200):
        base = {"identity_uuid": person.identity_uuid, "first_name": person.first_name,
                "last_name": person.last_name, "preferred_name": person.preferred_name,
                "email": person.contact_email, "phone": person.contact_phone,
                "date_of_birth": person.date_of_birth, "city": person.city,
                "state": person.state, "postal_code": person.postal_code}
        histories = list(person.historical_events.all())
        registrations = list(person.event_registrations.all())
        records = []
        for item in histories:
            base_record = {"event_name": item.event_name, "event_date": item.event_date,
                           "group_name": item.group_name, "registration_source": "historical",
                           "_event_id": item.event_id}
            related_scores = [score for score in person.historical_simulator_scores.all()
                              if score.event_record_id == item.pk]
            if related_scores:
                for score in related_scores:
                    records.append({**base_record, "simulator_score": score.raw_score,
                                    "score_method": score.score_method, "score_source": score.source})
            else:
                records.append(base_record)
        for item in registrations:
            records.append({"event_name": item.event.name, "event_date": item.event.start_at.date(),
                            "group_name": item.group.name if item.group else "",
                            "registration_source": item.source, "_event_id": item.event_id})
        for score in person.historical_simulator_scores.all():
            if score.event_record_id is None:
                records.append({"simulator_score": score.raw_score, "score_method": score.score_method,
                                "score_source": score.source})
        if not records:
            records = [{}]
        for record in records:
            if filters.get("event_id") and str(record.get("_event_id")) != str(event_id):
                continue
            if filters.get("group") and record.get("group_name", "").casefold() != filters["group"].casefold():
                continue
            if filters.get("registration_source") and record.get("registration_source") != source:
                continue
            record_date = record.get("event_date")
            if start_date and (record_date is None or record_date < start_date):
                continue
            if end_date and (record_date is None or record_date > end_date):
                continue
            values = {**base, **record}
            rows.append([safe_spreadsheet_cell(values.get(field)) for field in fields])
            if len(rows) > 50000:
                raise ValidationError("Export exceeds 50,000 rows. Narrow the filters.")
    if file_format == "csv":
        text = io.StringIO(newline="")
        writer = csv.writer(text)
        writer.writerow(fields)
        writer.writerows(rows)
        result = text.getvalue().encode("utf-8-sig")
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Participants"
        sheet.append(list(fields))
        for row in rows:
            sheet.append(row)
        for sheet_row in sheet.iter_rows():
            for cell in sheet_row:
                cell.data_type = "s"
        stream = io.BytesIO()
        workbook.save(stream)
        result = stream.getvalue()
    from .models import HistoricalExportRecord
    record = HistoricalExportRecord.objects.create(created_by=actor, criteria={"fields": list(fields),
                                                   "filters": filters}, row_count=len(rows), format=file_format)
    audit_action(actor=actor, action="historical.export", instance=record,
                 new_data={"rows": len(rows), "format": file_format})
    return result, len(rows)
